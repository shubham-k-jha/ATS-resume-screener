from io import BytesIO
import re

def parse_document(raw,name):
    if not raw: raise ValueError('The uploaded document is empty.')
    ext=name.lower().rsplit('.',1)[-1] if '.' in name else ''
    try:
        if ext=='txt': text=raw.decode('utf-8',errors='replace'); meta={'type':'TXT'}
        elif ext=='docx':
            from docx import Document
            d=Document(BytesIO(raw)); parts=[p.text for p in d.paragraphs if p.text.strip()]
            for table in d.tables:
                parts += [' | '.join(cell.text.strip() for cell in row.cells) for row in table.rows]
            text='\n'.join(parts); meta={'type':'DOCX','tables':len(d.tables),'paragraphs':len(d.paragraphs)}
        elif ext=='pdf':
            import fitz
            d=fitz.open(stream=raw,filetype='pdf'); pages=[p.get_text('text') for p in d]; text='\n'.join(pages)
            meta={'type':'PDF','pages':len(d),'images':sum(len(p.get_images(full=True)) for p in d),'page_text_lengths':[len(x.strip()) for x in pages]}
            if len(text.strip())<80: meta['scanned_warning']=True
        else: raise ValueError('Unsupported file type. Use PDF, DOCX, or TXT.')
    except Exception as e:
        if isinstance(e,ValueError): raise
        raise ValueError(f'Could not parse document: {e}')
    return text,meta

def clean_text(text):
    text=re.sub(r'[ \t]+',' ',text); text=re.sub(r'\n{3,}','\n\n',text); return text.strip()
